import sys
import openpyxl
import json

file_path = r'c:\Projetos\flutter\finance_control\CONTROLE FINANCEIRO_2026_R1.xlsx'

try:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    result = {}
    result['sheets'] = wb.sheetnames
    result['data'] = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_data = []
        max_cols = min(ws.max_column, 30)
        max_rows = min(ws.max_row, 200)
        
        for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
            rows_data.append([str(cell) if cell is not None else '' for cell in row])
        
        result['data'][sheet_name] = rows_data
    
    print(json.dumps(result, ensure_ascii=False, indent=2))
    
except Exception as e:
    print(f'ERROR: {e}', file=sys.stderr)
    import traceback
    traceback.print_exc()
